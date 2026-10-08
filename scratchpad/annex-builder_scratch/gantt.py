from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
from datetime import date, timedelta
D=date
wb=Workbook(); ws=wb.active; ws.title="Gantt"
F="Arial"
R="Robin"
rows=[
("G","Academic milestones"),
("M","Form B submitted",D(2026,10,11),D(2026,10,11),R,"None","Form B_Guo_Zi_Qiang_Robin.docx"),
("M","Interim Report (Form E1) submitted",D(2026,12,6),D(2026,12,6),R,"Test bench evaluation report; Tier 1 results","Form E1"),
("M","Interim Presentation",D(2027,1,31),D(2027,1,31),R,"Form E1; Beacon build progress","Interim Presentation"),
("M","Final Report (Form E2) submitted",D(2027,3,21),D(2027,3,21),R,"Deployed Beacon; test bench evaluation report","Form E2"),
("M","Capstone period ends",D(2027,3,31),D(2027,3,31),R,"Form E2","None (period boundary)"),
("M","Final Presentation (after the period)",D(2027,4,11),D(2027,4,11),R,"Form E2","Final Presentation"),
("T","Write Interim Report (Form E1)",D(2026,11,23),D(2026,12,6),R,"Tier 1 results","Form E1 draft"),
("T","Prepare Interim Presentation",D(2027,1,18),D(2027,1,31),R,"Beacon build under way","Slides and demo"),
("T","Write Final Report (Form E2)",D(2027,3,1),D(2027,3,21),R,"Beacon deployed and monitored","Form E2 draft"),
("T","Prepare Final Presentation",D(2027,3,22),D(2027,4,11),R,"Form E2","Slides and demo"),
("G","Workstream 1: Guardrail Evaluation Test Bench"),
("T","Complete comparative research (add Purple Llama and Litmus)",D(2026,10,8),D(2026,10,31),R,"None","Updated product research table"),
("T","Build minimal reusable evaluation harness",D(2026,10,12),D(2026,11,8),R,"Comparative research","Evaluation harness (code)"),
("T","Prepare labelled test inputs (public or synthetic only)",D(2026,10,26),D(2026,11,8),R,"Comparative research","Labelled test set"),
("T","Document Tier 2 items under investigation, with reasons",D(2026,11,2),D(2026,11,30),R,"Comparative research","Tier 2 reasons note"),
("T","Run Tier 1 evaluation (harmful-content and jailbreak checks)",D(2026,11,9),D(2026,11,27),R,"Harness; labelled test inputs","Raw results dataset"),
("T","Analyse results and compare guardrails",D(2026,11,16),D(2026,12,1),R,"Evaluation runs","Comparison tables"),
("M","Tier 1 evaluation results ready",D(2026,12,1),D(2026,12,1),R,"Evaluation runs","Results dataset"),
("T","Write test bench evaluation report",D(2026,12,1),D(2026,12,6),R,"Tier 1 results","Test bench evaluation report"),
("T","Tidy harness, dataset and documentation",D(2026,12,7),D(2026,12,18),R,"Test bench evaluation report","Documented code and dataset"),
("G","Workstream 2: Beacon"),
("T","Beacon design continues (low intensity)",D(2026,10,8),D(2026,11,30),R,"None","Design notes"),
("T","Build collection (sources, page reading, storage)",D(2026,12,1),D(2026,12,20),R,"Beacon design","Daily collection run"),
("T","Build filtering (on-topic check and labelling)",D(2026,12,14),D(2027,1,3),R,"Collection","Filtered, labelled items"),
("T","Build merging of duplicate reports",D(2026,12,28),D(2027,1,17),R,"Filtering","Grouped developments"),
("T","Build significance rating",D(2027,1,4),D(2027,1,24),R,"Merging","Rated developments"),
("T","Build web Briefing and Telegram digest",D(2027,1,11),D(2027,1,29),R,"Rating","Web Briefing; Telegram digest"),
("T","Initial mapping to SSP controls and MITRE ATT&CK / ATLAS",D(2027,1,18),D(2027,1,31),R,"Rating","Mapped developments with evidence"),
("T","Integrate and deploy",D(2027,1,25),D(2027,2,1),R,"Briefing, digest and mapping built","Deployed Beacon"),
("M","Beacon deployed",D(2027,2,1),D(2027,2,1),R,"Integration","Live service"),
("T","CI/CD pipeline",D(2027,2,1),D(2027,2,28),R,"Deployed Beacon","Automated build and release"),
("T","Monitoring and fixes",D(2027,2,1),D(2027,2,28),R,"Deployed Beacon","Monitoring and incident notes"),
("T","Trend detection (stretch goal)",D(2027,2,15),D(2027,3,15),R,"Stable pipeline","Trend view, if time allows"),
("T","Buffer and iteration",D(2027,3,1),D(2027,3,31),R,"Monitoring results","Fixes; source code and documentation"),
]
start=D(2026,10,8); nweeks=27; fc=8
hdr=5
ws["A1"]="Gantt chart: AI Security for Critical Infrastructure (Guardrail Evaluation Test Bench and Beacon)"; ws["A1"].font=Font(name=F,bold=True,size=14)
ws["A2"]="Guo Zi Qiang Robin, AAI4001 Capstone at LTA (CYAD). Capstone period 8 Oct 2026 to 31 Mar 2027. Weekly columns start on the date shown and run Thursday to Wednesday. See the Legend sheet."
ws["A2"].font=Font(name=F,italic=True,size=9)
ws["A3"]="Chart start date"; ws["B3"]=start; ws["B3"].number_format="d mmm yyyy"; ws["A3"].font=Font(name=F,bold=True,size=9); ws["B3"].font=Font(name=F,color="0000FF",size=9)
heads=["Task","Type","Start","End","Owner","Depends on","Deliverable"]
thin=Side(style="thin",color="D0D5DD")
for i,h in enumerate(heads,1):
    c=ws.cell(hdr,i,h); c.font=Font(name=F,bold=True,color="FFFFFF",size=10); c.fill=PatternFill("solid",fgColor="1F3A68"); c.alignment=Alignment(vertical="center",wrap_text=True)
ws.cell(hdr-1,fc,"Month").font=Font(name=F,bold=True,size=9)
for k in range(nweeks):
    col=fc+k
    c=ws.cell(hdr,col); c.value="=$B$3" if k==0 else f"={L(col-1)}{hdr}+7"
    c.number_format="d mmm"; c.font=Font(name=F,bold=True,color="FFFFFF",size=8)
    c.fill=PatternFill("solid",fgColor="1F3A68"); c.alignment=Alignment(text_rotation=90,horizontal="center",vertical="center")
    m=ws.cell(hdr-1,col); m.value=f'=TEXT({L(col)}{hdr},"mmm yy")' if k==0 else f'=IF(TEXT({L(col)}{hdr},"mmm yy")=TEXT({L(col-1)}{hdr},"mmm yy"),"",TEXT({L(col)}{hdr},"mmm yy"))'
    m.font=Font(name=F,bold=True,size=8)
    ws.column_dimensions[L(col)].width=3.6
ws.row_dimensions[hdr].height=42
r=hdr+1; first=r
for row in rows:
    if row[0]=="G":
        ws.cell(r,1,row[1])
        for col in range(1,fc+nweeks):
            c=ws.cell(r,col); c.fill=PatternFill("solid",fgColor="D9E2F3"); c.font=Font(name=F,bold=True,size=10)
        ws.cell(r,2,"Group")
    else:
        t,name,s,e,o,dep,dl=row
        for i,v in enumerate([name,{"M":"Milestone","T":"Task"}[t],s,e,o,dep,dl],1):
            c=ws.cell(r,i,v); c.font=Font(name=F,size=9,color=("0000FF" if i in(3,4) else "000000")); c.alignment=Alignment(wrap_text=True,vertical="center")
        ws.cell(r,3).number_format=ws.cell(r,4).number_format="d mmm yyyy"
        for k in range(nweeks):
            col=fc+k; cl=L(col)
            c=ws.cell(r,col,f'=IF(AND($B{r}="Milestone",$C{r}<={cl}${hdr}+6,$D{r}>={cl}${hdr}),"◆","")')
            c.alignment=Alignment(horizontal="center",vertical="center"); c.font=Font(name=F,size=9,bold=True,color="7A4B00")
    for col in range(1,fc+nweeks): ws.cell(r,col).border=Border(bottom=thin)
    r+=1
last=r-1
rng=f"{L(fc)}{first}:{L(fc+nweeks-1)}{last}"
ms=f'AND($B{first}="Milestone",$C{first}<={L(fc)}${hdr}+6,$D{first}>={L(fc)}${hdr})'
bar=f'AND($B{first}="Task",$C{first}<={L(fc)}${hdr}+6,$D{first}>={L(fc)}${hdr})'
stretch=f'AND($B{first}="Task",ISNUMBER(SEARCH("stretch",$A{first})),$C{first}<={L(fc)}${hdr}+6,$D{first}>={L(fc)}${hdr})'
ws.conditional_formatting.add(rng,FormulaRule(formula=[ms],fill=PatternFill("solid",bgColor="FFD966",fgColor="FFD966"),stopIfTrue=True))
ws.conditional_formatting.add(rng,FormulaRule(formula=[stretch],fill=PatternFill("solid",bgColor="A6A6A6",fgColor="A6A6A6"),stopIfTrue=True))
ws.conditional_formatting.add(rng,FormulaRule(formula=[bar],fill=PatternFill("solid",bgColor="2E75B6",fgColor="2E75B6")))
# today marker column: outline for the week containing 8 Oct not needed
for col,w in zip("ABCDEFG",[46,10,12,12,8,30,32]): ws.column_dimensions[col].width=w
ws.freeze_panes=ws.cell(hdr+1,fc)
ws.sheet_view.zoomScale=90
ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
ws.sheet_properties.pageSetUpPr.fitToPage=True
# legend sheet
lg=wb.create_sheet("Legend")
items=[("Blue bar","Task: shaded in every week the task is active (formula-driven from Start and End)","2E75B6"),
("Amber cell with a diamond","Milestone: a fixed date set by SIT or by the plan","FFD966"),
("Grey bar","Stretch goal: done only if earlier stages are stable","A6A6A6"),
("Light-blue row","Group heading","D9E2F3")]
lg["A1"]="Legend"; lg["A1"].font=Font(name=F,bold=True,size=14)
for i,(a,b,c) in enumerate(items,3):
    lg.cell(i,1,a).fill=PatternFill("solid",fgColor=c); lg.cell(i,1).font=Font(name=F,bold=True,size=10); lg.cell(i,2,b).font=Font(name=F,size=10)
notes=["How to edit: change Start or End dates (blue text, columns C and D) and the bars redraw. Change the chart start date in Gantt!B3 to shift the weekly columns.",
"Weekly columns: 27 weeks from Thu 8 Oct 2026; the last column (week of 8 Apr 2027) contains the Final Presentation on 11 Apr 2027.",
"Capstone period: 8 Oct 2026 to 31 Mar 2027. The Final Presentation on 11 Apr 2027 falls after the period.",
"Phasing: Workstream 1 is heavier from Oct to early Dec; Workstream 2 design continues at low intensity, then the build runs Dec to Jan.",
"Assumptions: single owner (the student) for all rows; sub-task dates inside each phase are planning estimates; Trend detection is a stretch goal; Tier 2 comparisons are under investigation and not scheduled for evaluation.",
"Example of a row: 'Run Tier 1 evaluation' | Task | 9 Nov 2026 | 27 Nov 2026 | Robin | Harness; labelled test inputs | Raw results dataset."]
for i,n in enumerate(notes,9):
    lg.cell(i,1,n).font=Font(name=F,size=10); lg.merge_cells(start_row=i,start_column=1,end_row=i,end_column=2); lg.cell(i,1).alignment=Alignment(wrap_text=True,vertical="top"); lg.row_dimensions[i].height=32
lg.column_dimensions["A"].width=30; lg.column_dimensions["B"].width=100
wb.save("Gantt_Guo_Zi_Qiang_Robin.xlsx")
